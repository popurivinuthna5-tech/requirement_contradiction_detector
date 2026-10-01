import docx
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib import colors

def create_sample_docx(txt_path: Path, docx_path: Path, title: str):
    doc = docx.Document()
    doc.add_heading(title, 0)
    
    with open(txt_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    for line in lines:
        line_s = line.strip()
        if not line_s or line_s.startswith("#"):
            continue
        if line_s.startswith("SECTION"):
            doc.add_heading(line_s, level=1)
        elif line_s[0].isdigit() and "." in line_s[:4]:
            p = doc.add_paragraph()
            parts = line_s.split(" ", 1)
            p.add_run(parts[0] + " ").bold = True
            if len(parts) > 1:
                p.add_run(parts[1])
        else:
            doc.add_paragraph(line_s)
            
    doc.save(str(docx_path))
    print(f"Generated DOCX: {docx_path}")

def create_sample_pdf(txt_path: Path, pdf_path: Path, title: str):
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    width, height = letter
    y = height - 50
    
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, y, title)
    y -= 30
    
    with open(txt_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    c.setFont("Helvetica", 10)
    for line in lines:
        line_s = line.strip()
        if not line_s or line_s.startswith("#"):
            continue
            
        if line_s.startswith("SECTION"):
            y -= 15
            c.setFont("Helvetica-Bold", 12)
            c.setFillColor(colors.HexColor("#1e293b"))
            c.drawString(50, y, line_s)
            c.setFont("Helvetica", 10)
            c.setFillColor(colors.black)
            y -= 18
        else:
            # Wrap long line
            words = line_s.split()
            current_line = []
            for w in words:
                current_line.append(w)
                if len(" ".join(current_line)) > 90:
                    c.drawString(60, y, " ".join(current_line[:-1]))
                    y -= 14
                    current_line = [w]
            if current_line:
                c.drawString(60, y, " ".join(current_line))
                y -= 16
                
        if y < 60:
            c.showPage()
            c.setFont("Helvetica", 10)
            y = height - 50
            
    c.save()
    print(f"Generated PDF: {pdf_path}")

if __name__ == "__main__":
    base = Path(__file__).parent
    srs_txt = base / "E-Commerce_Platform_SRS.txt"
    brd_txt = base / "E-Commerce_Business_Requirements_BRD.txt"
    
    create_sample_docx(srs_txt, base / "E-Commerce_Platform_SRS.docx", "Software Requirements Specification (SRS)")
    create_sample_docx(brd_txt, base / "E-Commerce_Business_Requirements_BRD.docx", "Business Requirements Document (BRD)")
    
    create_sample_pdf(srs_txt, base / "E-Commerce_Platform_SRS.pdf", "Software Requirements Specification (SRS)")
    create_sample_pdf(brd_txt, base / "E-Commerce_Business_Requirements_BRD.pdf", "Business Requirements Document (BRD)")
