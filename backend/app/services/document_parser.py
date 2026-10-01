import re
from pathlib import Path
from typing import List, Dict, Any, Optional
import pypdf
import docx

class DocumentBlock:
    def __init__(self, text: str, page_number: Optional[int] = None, section: Optional[str] = None):
        self.text = text.strip()
        self.page_number = page_number
        self.section = section or "General"

    def to_dict(self) -> Dict[str, Any]:
        return {
            "text": self.text,
            "page_number": self.page_number,
            "section": self.section
        }

class DocumentParser:
    """Parses PDF, DOCX, and TXT files, preserving structure, sections, and page numbers."""
    
    @classmethod
    def parse_file(cls, file_path: Path, document_type: str) -> List[DocumentBlock]:
        ext = file_path.suffix.lower()
        if ext == ".pdf":
            return cls.parse_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return cls.parse_docx(file_path)
        elif ext == ".txt":
            return cls.parse_txt(file_path)
        else:
            raise ValueError(f"Unsupported file format: {ext}. Supported formats: PDF, DOCX, TXT.")

    @classmethod
    def parse_pdf(cls, file_path: Path) -> List[DocumentBlock]:
        blocks: List[DocumentBlock] = []
        try:
            reader = pypdf.PdfReader(str(file_path))
            current_section = "General"
            
            for page_idx, page in enumerate(reader.pages):
                page_num = page_idx + 1
                page_text = page.extract_text() or ""
                lines = page_text.splitlines()
                
                buffer = []
                for line in lines:
                    line_clean = line.strip()
                    if not line_clean:
                        if buffer:
                            combined = " ".join(buffer)
                            if len(combined) > 10:
                                blocks.append(DocumentBlock(combined, page_number=page_num, section=current_section))
                            buffer = []
                        continue
                    
                    # Detect section headers (e.g. "Section 4.1 Security", "4.1 Authentication", "Chapter 2")
                    if cls._is_section_header(line_clean):
                        if buffer:
                            blocks.append(DocumentBlock(" ".join(buffer), page_number=page_num, section=current_section))
                            buffer = []
                        current_section = line_clean
                    else:
                        buffer.append(line_clean)
                
                if buffer:
                    blocks.append(DocumentBlock(" ".join(buffer), page_number=page_num, section=current_section))
                    
        except Exception as e:
            raise RuntimeError(f"Failed to parse PDF file '{file_path.name}': {str(e)}")
            
        return blocks

    @classmethod
    def parse_docx(cls, file_path: Path) -> List[DocumentBlock]:
        blocks: List[DocumentBlock] = []
        try:
            doc = docx.Document(str(file_path))
            current_section = "General"
            estimated_page = 1
            paragraph_count = 0
            
            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                
                # Check heading styles
                if p.style.name.startswith("Heading") or cls._is_section_header(text):
                    current_section = text
                    continue
                
                paragraph_count += 1
                # Rough page estimate for DOCX (every ~6-8 paragraphs is roughly a page if not paginated)
                estimated_page = (paragraph_count // 7) + 1
                
                blocks.append(DocumentBlock(text, page_number=estimated_page, section=current_section))
                
            # Also parse any tables in DOCX
            for table in doc.tables:
                for row in table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        combined = " | ".join(row_texts)
                        if len(combined) > 15:
                            blocks.append(DocumentBlock(combined, page_number=estimated_page, section=f"{current_section} (Table)"))
                            
        except Exception as e:
            raise RuntimeError(f"Failed to parse DOCX file '{file_path.name}': {str(e)}")
            
        return blocks

    @classmethod
    def parse_txt(cls, file_path: Path) -> List[DocumentBlock]:
        blocks: List[DocumentBlock] = []
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                lines = f.readlines()
                
            current_section = "General"
            buffer = []
            line_count = 0
            
            for line in lines:
                line_count += 1
                estimated_page = (line_count // 45) + 1 # Approximate page in standard manuscript
                line_clean = line.strip()
                
                if not line_clean:
                    if buffer:
                        combined = " ".join(buffer)
                        if len(combined) > 10:
                            blocks.append(DocumentBlock(combined, page_number=estimated_page, section=current_section))
                        buffer = []
                    continue
                
                # Detect section header
                if cls._is_section_header(line_clean):
                    if buffer:
                        blocks.append(DocumentBlock(" ".join(buffer), page_number=estimated_page, section=current_section))
                        buffer = []
                    current_section = line_clean
                else:
                    buffer.append(line_clean)
                    
            if buffer:
                blocks.append(DocumentBlock(" ".join(buffer), page_number=(line_count // 45) + 1, section=current_section))
                
        except Exception as e:
            raise RuntimeError(f"Failed to parse text file '{file_path.name}': {str(e)}")
            
        return blocks

    @staticmethod
    def _is_section_header(line: str) -> bool:
        """Determines if a given text line represents a section or chapter header."""
        if len(line) > 75 or len(line) < 3:
            return False
        
        # Sentences with verbs or ending with punctuation are requirements, not section headers
        line_clean = line.strip()
        if re.search(r'\b(must|shall|should|will|can|may|cannot|allows|provides)\b', line_clean, re.IGNORECASE):
            return False
        if line_clean.endswith(".") and not re.match(r'^\d+\.', line_clean):
            return False
            
        if line_clean.isupper() and len(line_clean) < 55:
            return True
            
        patterns = [
            r"^(?:section|chapter|part)\s+\d+",
            r"^\d+\.\d+(?:\.\d+)*\s+[A-Z][a-zA-Z\s\-]{2,35}$",
            r"^(?:functional|non-functional|security|performance|business|system)\s+requirements"
        ]
        for p in patterns:
            if re.search(p, line_clean, re.IGNORECASE):
                return True
        return False
